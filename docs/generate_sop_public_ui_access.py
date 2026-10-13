"""Generate the SOP .docx for enabling public UI/API access on a private Azure Databricks workspace.

Usage: pip install python-docx && python docs/generate_sop_public_ui_access.py
"""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT = Path(__file__).with_name("SOP_Enable_Public_UI_Access_Azure_Databricks.docx")
ACCENT = RGBColor(0x1F, 0x3A, 0x5F)

def shade(cell_or_par, fill):
    el = cell_or_par._tc if hasattr(cell_or_par, "_tc") else cell_or_par._p
    pr = el.get_or_add_tcPr() if hasattr(cell_or_par, "_tc") else el.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    pr.append(shd)


def para(doc, text="", bold_lead=None, size=None, space_after=4, style=None):
    p = doc.add_paragraph(style=style)
    if bold_lead:
        p.add_run(bold_lead).bold = True
    r = p.add_run(text)
    if size:
        r.font.size = Pt(size)
    p.paragraph_format.space_after = Pt(space_after)
    return p


def bullet(doc, text, bold_lead=None):
    p = para(doc, text, bold_lead, style="List Bullet", space_after=2)
    return p


def step(doc, num, title, lines):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(f"Step {num}. {title}")
    r.bold = True
    r.font.color.rgb = ACCENT
    for i, line in enumerate(lines, 1):
        q = doc.add_paragraph()
        q.paragraph_format.left_indent = Cm(1.0)
        q.paragraph_format.first_line_indent = Cm(-0.5)
        q.paragraph_format.space_after = Pt(1)
        q.add_run(f"{num}.{i}  ")
        if isinstance(line, tuple):
            q.add_run(line[0])
            q.add_run(line[1]).bold = True
            q.add_run(line[2])
        else:
            q.add_run(line)


def screenshot(doc, desc):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.5)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(4)
    shade(p, "E7E6E6")
    r = p.add_run(f"[SCREENSHOT PLACEHOLDER: {desc}]")
    r.italic = True
    r.font.size = Pt(9.5)
    r.font.color.rgb = RGBColor(0x40, 0x40, 0x40)


def build():
    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    sec.top_margin = sec.bottom_margin = Cm(2.0)
    sec.left_margin = sec.right_margin = Cm(2.2)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    for name, size in (("Heading 1", 14), ("Heading 2", 12)):
        st = doc.styles[name]
        st.font.name = "Calibri"
        st.font.size = Pt(size)
        st.font.color.rgb = ACCENT
        st.paragraph_format.space_before = Pt(10)
        st.paragraph_format.space_after = Pt(4)

    # Title / header block
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = title.add_run("SOP: Enabling Public UI Access for a Private Azure Databricks Workspace")
    r.bold = True
    r.font.size = Pt(18)
    r.font.color.rgb = ACCENT
    # 1. Purpose & Scope
    doc.add_heading("1. Purpose and Scope", level=1)
    para(doc, "This SOP describes how to allow users to reach the Azure Databricks workspace "
              "web application (UI) and workspace REST APIs over the public internet, for a workspace that is "
              "currently private, while keeping all classic compute (clusters, SQL warehouses on classic compute) "
              "and the data plane private.")
    bullet(doc, "Front-end (user → workspace) connectivity: UI login and REST API reachability from public networks.",
           "In scope: ")
    bullet(doc, "Secure Cluster Connectivity (No Public IP), VNet injection, subnets/NSGs, back-end Private Link, "
                "storage/data-access private endpoints. These are not changed by this SOP.", "Out of scope: ")
    para(doc, "All steps are performed in the Azure Portal and the Azure Databricks web UI. No CLI, ARM, Bicep or "
              "Terraform is used.")

    # 2. Procedure
    doc.add_heading("2. Procedure", level=1)
    step(doc, 1, "Sign in and open the workspace resource", [
        ("Sign in to ", "https://portal.azure.com", " with an account that has Contributor/Owner on the workspace."),
        "Search for “Azure Databricks”, open the service and select the target workspace.",
        "On Overview, record the Workspace URL (adb-<id>.<n>.azuredatabricks.net) and resource group.",
    ])
    screenshot(doc, "Azure Databricks workspace Overview blade showing workspace name, resource group and Workspace URL")

    step(doc, 2, "Open the Networking settings and record the current state", [
        ("In the left menu, under ", "Settings", ", select Networking."),
        "Record the current values (needed to revert the change): Allow Public Network Access (expected: Disabled), "
        "Required NSG rules (e.g. NoAzureDatabricksRules), Secure cluster connectivity (No Public IP) (expected: Enabled).",
        ("Open the ", "Private endpoint connections", " tab and note the existing endpoints and their Approved status."),
    ])
    screenshot(doc, "Networking blade – current settings with Allow Public Network Access = Disabled")
    screenshot(doc, "Private endpoint connections tab listing existing private endpoints (before change)")

    step(doc, 3, "Enable public network access (front-end only)", [
        ("Set ", "Allow Public Network Access", " to Enabled."),
        ("Leave ", "Required NSG rules", " and all other networking settings unchanged."),
        ("Click ", "Save", " and wait for the update notification to report success (may take several minutes)."),
    ])
    screenshot(doc, "Allow Public Network Access set to Enabled, Required NSG rules unchanged, before clicking Save")
    screenshot(doc, "Azure notification showing the workspace update succeeded")

    step(doc, 4, "(Recommended) Restrict public access to approved IP ranges", [
        ("Sign in to the Azure Databricks account console (", "https://accounts.azuredatabricks.net", ") as an account admin."),
        ("Go to ", "Security → Networking → Context-based ingress & egress control", " and click Create new network policy."),
        "Enter a policy name. On the Ingress tab, set public network access to “Restrict access based on request "
        "context”. Add Allow rules for Workspace UI and API, identity “All users and service principals”, "
        "network source “Selected IPs”, and enter the approved public IP/CIDR ranges.",
        "Select Dry run mode for all products first, click Create, and review denials in "
        "system.access.inbound_network; switch to Enforced for all products once validated.",
        ("Go to ", "Workspaces", " → select the workspace → Network Policy → Update network policy → choose the "
         "policy → Apply policy. Allow 10–15 minutes to take effect."),
    ])
    screenshot(doc, "Account console – new network policy Ingress tab with Selected IPs allow rules for Workspace UI and API")
    screenshot(doc, "Account console – workspace page showing the network policy attached")

    step(doc, 5, "Verify UI access and that compute is still private", [
        "From a device on an allowed public network (not on VPN/private DNS), open the Workspace URL and confirm the "
        "Microsoft Entra ID sign-in completes and the workspace home page loads.",
        "If an IP restriction is enforced, confirm access from a non-allowed public IP is denied.",
        "In the Azure Portal Networking blade, confirm Secure cluster connectivity (No Public IP) is still Enabled and "
        "Required NSG rules and private endpoints are unchanged.",
        "Start a test cluster. In the workspace’s managed resource group, open a cluster VM → Networking and confirm "
        "it has no public IP address. Run a simple notebook command to confirm the data plane works.",
    ])
    screenshot(doc, "Workspace UI home page loaded from a public network (browser address bar visible)")
    screenshot(doc, "Cluster VM network interface in the managed resource group showing no public IP")

    doc.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    build()
